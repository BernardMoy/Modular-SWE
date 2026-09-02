"""Black-box tests for checkpoint 3 formula-aware rule kinds.

The implementation under test is selected with SHEETEVAL_IMPLEMENTATION and
defaults to the repository's previous implementation.  Workbooks are made in
tmp_path so these tests do not rely on a fixture dataset or a fixed path.
"""

from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys

import openpyxl
import pytest
from openpyxl.comments import Comment


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def implementation_dir() -> Path:
    configured = os.environ.get("SHEETEVAL_IMPLEMENTATION")
    return Path(configured) if configured else ROOT / "previous_implementation"


def make_workbooks(tmp_path: Path, *, answer_value, submission_value,
                   rule: str, cell: str = "B2", flags: str | None = None,
                   answer_alternates: dict[str, object] | None = None,
                   answer_formula_cells: dict[str, str] | None = None,
                   submission_cells: dict[str, object] | None = None):
    """Create the smallest workbook pair observable by the CLI."""
    answer = openpyxl.Workbook()
    answer.remove(answer.active)
    manifest = answer.create_sheet("EvalManifest")
    manifest.append(["name"])
    manifest.append(["Sheet1"])
    sheet = answer.create_sheet("Sheet1")
    sheet[cell] = answer_value
    sheet[cell].comment = Comment(rule, "test")
    checklist = answer.create_sheet("Sheet1_Checklist")
    checklist.append(["number", "flags"])
    checklist.append([cell, flags])
    for ref, value in (answer_alternates or {}).items():
        sheet[ref] = value
    for ref, formula in (answer_formula_cells or {}).items():
        sheet[ref] = formula

    submission = openpyxl.Workbook()
    sub_sheet = submission.active
    sub_sheet.title = "Sheet1"
    sub_sheet[cell] = submission_value
    for ref, value in (submission_cells or {}).items():
        sub_sheet[ref] = value

    answer_path = tmp_path / "answer.xlsx"
    submission_path = tmp_path / "submission.xlsx"
    answer.save(answer_path)
    submission.save(submission_path)
    return answer_path, submission_path


def run_cli(implementation_dir: Path, answer: Path, submission: Path,
            *options: str, timeout: float = 5) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(implementation_dir / "sheeteval.py"),
         str(answer), str(submission), *options],
        cwd=implementation_dir, text=True, capture_output=True, timeout=timeout,
    )


def grade(tmp_path, implementation_dir, *, answer_value, submission_value,
          rule, **kwargs):
    options = kwargs.pop("options", ())
    timeout = kwargs.pop("timeout", 5)
    answer, submission = make_workbooks(
        tmp_path, answer_value=answer_value, submission_value=submission_value,
        rule=rule, **kwargs)
    return run_cli(implementation_dir, answer, submission, *options, timeout=timeout)


@pytest.mark.parametrize(
    ("expected", "actual"),
    [
        ("=SUM(A1:A2)", "=sum( A1 : A2 )"),
        ("=A1 + 2", "= a1   +   2"),
        ("=A1+B1", "=A1-B1"),  # structural token differences remain unequal
    ],
)
def test_expression_normalization_and_structural_matching(
    tmp_path, implementation_dir, expected, actual
):
    result = grade(
        tmp_path, implementation_dir, answer_value=expected,
        submission_value=actual, rule="rule:\n  kind: expression\n  points: 2",
    )
    if expected == "=A1+B1":
        assert "Assignment Score: 0 / 2" in result.stdout
    else:
        assert "Assignment Score: 2 / 2" in result.stdout


def test_expression_rejects_plain_submission_value(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value="=A1+1",
                   submission_value=3,
                   rule="rule:\n  kind: expression\n  points: 2")
    assert "Assignment Score: 0 / 2" in result.stdout
    assert result.stderr == ""


def test_expression_requires_formula_in_answer_key(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value="A1+1",
                   submission_value="=A1+1",
                   rule="rule:\n  kind: expression\n  points: 1")
    assert "Assignment Score: 0 / 1" in result.stdout
    assert "Sheet1.B2" in result.stderr
    assert "formula" in result.stderr.lower()


@pytest.mark.parametrize(
    "formula, cells, expected",
    [
        ("=A1+B1*2", {"A1": 3, "B1": 4}, 11),
        ("=-(A1^2)+B1/2", {"A1": 3, "B1": 8}, -5),
        ('="left"&"right"', {}, "leftright"),
        ("=A1>3", {"A1": 4}, 1),
        ("=A1=3", {"A1": 4}, 0),
    ],
)
def test_computed_supports_references_precedence_comparisons_and_concat(
    tmp_path, implementation_dir, formula, cells, expected
):
    result = grade(
        tmp_path, implementation_dir, answer_value=expected,
        submission_value=formula,
        submission_cells=cells,
        rule="rule:\n  kind: computed\n  points: 1",
    )
    assert "Assignment Score: 1 / 1" in result.stdout


def test_computed_references_and_functions_are_case_insensitive(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value=7,
                   submission_value="=sUm(bc42, a1)",
                   submission_cells={"BC42": 4, "A1": 3},
                   rule="rule:\n  kind: computed\n  points: 1")
    assert "Assignment Score: 1 / 1" in result.stdout


@pytest.mark.parametrize(
    "formula, cells, expected",
    [
        ("=SUM(D1:F2)", {"D1": 1, "E1": 2, "F1": "x", "D2": 3, "E2": 4, "F2": 5}, 15),
        ("=AVERAGE(A1:C1)", {"A1": 2, "B1": "x", "C1": 4}, 3),
        ("=MIN(A1:C1)", {"A1": 2, "B1": 7, "C1": 4}, 2),
        ("=MAX(A1:C1)", {"A1": 2, "B1": 7, "C1": 4}, 7),
        ("=COUNT(A1:C1)", {"A1": 2, "B1": "x", "C1": 4}, 2),
        ('=COUNTIF(A1:A4,">5")', {"A1": 6, "A2": 5, "A3": 8, "A4": 0}, 2),
        ('=COUNTIF(A1:A4,"<>0")', {"A1": 6, "A2": 0, "A3": "x", "A4": 2}, 3),
        ('=COUNTIF(A1:A4,"yes")', {"A1": "yes", "A2": "no", "A3": "yes", "A4": 1}, 2),
        ("=ABS(A1)", {"A1": -4}, 4),
        ("=ROUND(A1,2)", {"A1": 1.236}, 1.24),
        ("=INT(A1)", {"A1": 3.9}, 3),
        ("=MOD(A1,B1)", {"A1": 8, "B1": 3}, 2),
        ('=IF(A1>1,"yes","no")', {"A1": 2}, "yes"),
    ],
)
def test_computed_supported_functions(tmp_path, implementation_dir, formula,
                                      cells, expected):
    result = grade(tmp_path, implementation_dir, answer_value=expected,
                   submission_value=formula, submission_cells=cells,
                   rule="rule:\n  kind: computed\n  points: 1")
    assert "Assignment Score: 1 / 1" in result.stdout


def test_computed_requires_formula_and_uses_tolerance(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value=10,
                   submission_value=10.05,
                   rule="rule:\n  kind: computed\n  points: 3\n  tolerance: 0.1")
    assert "Assignment Score: 0 / 3" in result.stdout
    result = grade(tmp_path, implementation_dir, answer_value=10,
                   submission_value="=9.95",
                   rule="rule:\n  kind: computed\n  points: 3\n  tolerance: 0.1")
    assert "Assignment Score: 3 / 3" in result.stdout


def test_computed_rejects_plain_submission_value_without_warning(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value=10,
                   submission_value=10,
                   rule="rule:\n  kind: computed\n  points: 3")
    assert "Assignment Score: 0 / 3" in result.stdout
    assert result.stderr == ""


def test_expression_tolerance_does_not_change_string_comparison(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value="=A1+1",
                   submission_value="=A1+2",
                   rule="rule:\n  kind: expression\n  points: 2\n  tolerance: 100")
    assert "Assignment Score: 0 / 2" in result.stdout


def test_alternates_are_first_match_for_expression_and_computed(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value="=A1+1",
                   submission_value="=B1+2",
                   answer_alternates={"C2": "=B1+2"},
                   rule="rule:\n  kind: expression\n  points: 1\n  alternates: [C2]")
    assert "Assignment Score: 1 / 1" in result.stdout
    result = grade(tmp_path, implementation_dir, answer_value=10,
                   submission_value="=9",
                   answer_alternates={"C2": 9},
                   rule="rule:\n  kind: computed\n  points: 1\n  alternates: [C2]")
    assert "Assignment Score: 1 / 1" in result.stdout


def test_penalty_mode_scores_zero_for_correct_and_negative_for_incorrect(
    tmp_path, implementation_dir
):
    result = grade(tmp_path, implementation_dir, answer_value=5,
                   submission_value="=5",
                   rule="rule:\n  kind: computed\n  points: 4\n  mode: penalty")
    assert "Assignment Score: 0 / 0" in result.stdout
    result = grade(tmp_path, implementation_dir, answer_value=5,
                   submission_value="=6",
                   rule="rule:\n  kind: computed\n  points: 4\n  mode: penalty")
    assert "Assignment Score: -4 / 0" in result.stdout


@pytest.mark.parametrize("mode", ["award", "penalty"])
def test_formula_errors_score_zero_and_warn(tmp_path, implementation_dir, mode):
    result = grade(tmp_path, implementation_dir, answer_value=1,
                   submission_value="=UNKNOWNFUNC(1)",
                   rule=f"rule:\n  kind: computed\n  points: 2\n  mode: {mode}")
    assert result.returncode == 0
    assert "Assignment Score: 0 /" in result.stdout
    assert "UNKNOWNFUNC" in result.stderr
    assert "Sheet1.B2" in result.stderr


def test_circular_reference_warns_and_does_not_hang(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value=1,
                   submission_value="=A1+1", submission_cells={"A1": "=B1", "B1": "=A1"},
                   rule="rule:\n  kind: computed\n  points: 1", timeout=2)
    assert result.returncode == 0
    assert "Assignment Score: 0 /" in result.stdout
    assert "circular" in result.stderr.lower()


def test_malformed_formula_warns_and_scores_zero(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value=1,
                   submission_value="=SUM(A1:",
                   rule="rule:\n  kind: computed\n  points: 1")
    assert result.returncode == 0
    assert "Assignment Score: 0 /" in result.stdout
    assert "Sheet1.B2" in result.stderr


def test_quiet_warnings_suppresses_formula_warning(tmp_path, implementation_dir):
    result = grade(tmp_path, implementation_dir, answer_value=1,
                   submission_value="=NOPE(1)",
                   rule="rule:\n  kind: computed\n  points: 1", options=("--quiet-warnings",))
    # Keep invocation black-box while allowing the option to be passed through.
    assert "Assignment Score: 0 /" in result.stdout
    assert result.stderr == ""
