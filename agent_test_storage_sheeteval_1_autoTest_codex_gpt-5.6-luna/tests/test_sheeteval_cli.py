"""Black-box tests for the sheeteval command-line contract."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment


ROOT = Path(__file__).resolve().parents[1]
ENTRYPOINT = ROOT / "implementation" / "sheeteval.py"


def make_workbook(path: Path, sheets: dict[str, dict[str, object]]) -> Path:
    """Create a workbook from {sheet: {cell: value}} mappings."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, cells in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for coordinate, value in cells.items():
            sheet[coordinate] = value
    workbook.save(path)
    return path


def make_answer_key(
    path: Path,
    *,
    manifest_headers: list[str] | None = None,
    manifest_rows: list[list[object]] | None = None,
    checklist_headers: list[str] | None = None,
    checklist_rows: list[list[object]] | None = None,
    values: dict[str, dict[str, object]] | None = None,
    annotations: dict[tuple[str, str], str | None] | None = None,
    checklist_flags: dict[str, object] | None = None,
) -> Path:
    manifest_headers = ["name"] if manifest_headers is None else manifest_headers
    manifest_rows = [["Sheet1"]] if manifest_rows is None else manifest_rows
    checklist_headers = ["number"] if checklist_headers is None else checklist_headers
    checklist_rows = [["A1"]] if checklist_rows is None else checklist_rows
    values = {"Sheet1": {"A1": 1}} if values is None else values
    annotations = {("Sheet1", "A1"): "rule:\n  kind: literal\n  points: 1"} if annotations is None else annotations
    checklist_flags = {} if checklist_flags is None else checklist_flags

    workbook = Workbook()
    workbook.remove(workbook.active)
    manifest = workbook.create_sheet("EvalManifest")
    manifest.append(manifest_headers)
    for row in manifest_rows:
        manifest.append(row)

    for sheet_name, cells in values.items():
        sheet = workbook.create_sheet(sheet_name)
        for coordinate, value in cells.items():
            sheet[coordinate] = value
            annotation = annotations.get((sheet_name, coordinate))
            if annotation is not None:
                sheet[coordinate].comment = Comment(annotation, "test")

    for sheet_name in manifest_rows_to_names(manifest_rows, manifest_headers):
        checklist = workbook.create_sheet(f"{sheet_name}_Checklist")
        checklist.append(checklist_headers)
        for row in checklist_rows:
            checklist.append(row + ([checklist_flags.get(str(row[0]), "")] if "flags" in checklist_headers else []))

    workbook.save(path)
    return path


def manifest_rows_to_names(rows: list[list[object]], headers: list[str]) -> list[str]:
    index = next(i for i, header in enumerate(headers) if header in {"name", "sheet", "tab"})
    return [str(row[index]) for row in rows]


def make_submission(path: Path, values: dict[str, dict[str, object]]) -> Path:
    return make_workbook(path, values)


def run_cli(answer_key: Path, submission: Path, *options: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(ENTRYPOINT), str(answer_key), str(submission), *options],
        text=True,
        capture_output=True,
    )


@pytest.fixture
def basic_files(tmp_path: Path) -> tuple[Path, Path]:
    answer = make_answer_key(tmp_path / "answer.xlsx")
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    return answer, submission


def test_literal_scoring_and_exact_stdout(basic_files: tuple[Path, Path]) -> None:
    answer, submission = basic_files
    result = run_cli(answer, submission)
    assert result.returncode == 0
    assert result.stdout == "Assignment Score: 1 / 1\n"
    assert result.stderr == ""


def test_empty_manifest_has_a_zero_score(tmp_path: Path) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx", manifest_rows=[], values={})
    submission = make_submission(tmp_path / "submission.xlsx", {})
    result = run_cli(answer, submission)
    assert result.returncode == 0
    assert result.stdout == "Assignment Score: 0 / 0\n"
    assert result.stderr == ""


def test_type_sensitive_literal_comparison(basic_files: tuple[Path, Path], tmp_path: Path) -> None:
    answer, _ = basic_files
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": "1"}})
    result = run_cli(answer, submission)
    assert result.returncode == 0
    assert result.stdout == "Assignment Score: 0 / 1\n"


def test_score_file_uses_numeric_rendering(tmp_path: Path) -> None:
    answer = make_answer_key(
        tmp_path / "answer.xlsx",
        checklist_rows=[["A1"], ["B1"], ["C1"]],
        values={"Sheet1": {"A1": 1, "B1": 2, "C1": 3}},
        annotations={
            ("Sheet1", "A1"): "rule:\n  kind: literal\n  points: 3.50",
            ("Sheet1", "B1"): "rule:\n  kind: literal\n  points: 2.75",
            ("Sheet1", "C1"): "rule:\n  kind: literal\n  points: 1",
        },
    )
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1, "B1": 0, "C1": 3}})
    score_file = tmp_path / "score.txt"
    result = run_cli(answer, submission, "--score-file", str(score_file))
    assert result.returncode == 0
    assert result.stdout == "Assignment Score: 4.5 / 7.25\n"
    assert score_file.read_text() == "4.5\n"


def test_manifest_and_checklist_aliases_leftmost_and_order_are_honored(tmp_path: Path) -> None:
    answer = make_answer_key(
        tmp_path / "answer.xlsx",
        manifest_headers=["name", "sheet", "ignored"],
        manifest_rows=[["First", "Wrong", "x"], ["Second", "Wrong2", "y"]],
        checklist_headers=["number", "cell-id", "flags", "ignored"],
        checklist_rows=[["A1", "B1"], ["B1", "A1"]],
        values={"First": {"A1": "a", "B1": "b"}, "Second": {"A1": "c", "B1": "d"}},
        annotations={
            ("First", "A1"): "rule:\n  kind: literal\n  points: 1",
            ("First", "B1"): "rule:\n  kind: literal\n  points: 2",
            ("Second", "A1"): "rule:\n  kind: literal\n  points: 4",
            ("Second", "B1"): "rule:\n  kind: literal\n  points: 8",
        },
    )
    submission = make_submission(tmp_path / "submission.xlsx", {"First": {"A1": "a", "B1": "b"}, "Second": {"A1": "c", "B1": "d"}})
    result = run_cli(answer, submission, "--log")
    assert result.returncode == 0
    assert result.stdout == "Assignment Score: 15 / 15\n"
    assert result.stderr.splitlines() == [
        "LOG First.A1: literal CORRECT",
        "LOG First.B1: literal CORRECT",
        "LOG Second.A1: literal CORRECT",
        "LOG Second.B1: literal CORRECT",
    ]


def test_duplicate_coordinates_are_evaluated_independently(tmp_path: Path) -> None:
    answer = make_answer_key(
        tmp_path / "answer.xlsx",
        checklist_rows=[["A1"], ["A1"]],
        annotations={("Sheet1", "A1"): "rule:\n  kind: literal\n  points: 2"},
    )
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    result = run_cli(answer, submission, "--log")
    assert result.stdout == "Assignment Score: 4 / 4\n"
    assert result.stderr.splitlines() == ["LOG Sheet1.A1: literal CORRECT"] * 2


@pytest.mark.parametrize("flags", ["manual", "MANUAL", "Manual"])
def test_manual_rules_are_excluded_and_not_logged(tmp_path: Path, flags: str) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx", checklist_headers=["number", "flags"], checklist_rows=[["A1", flags]])
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    result = run_cli(answer, submission, "--log")
    assert result.stdout == "Assignment Score: 0 / 0\n"
    assert result.stderr == ""


def test_missing_checklist_is_a_suppressible_warning_and_skips_sheet(tmp_path: Path) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx")
    # Re-save without the generated checklist while preserving the manifest and graded sheet.
    workbook = load_workbook(answer)
    del workbook["Sheet1_Checklist"]
    workbook.save(answer)
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    result = run_cli(answer, submission)
    assert result.stdout == "Assignment Score: 0 / 0\n"
    assert "Sheet1_Checklist" in result.stderr
    quiet = run_cli(answer, submission, "--quiet-warnings")
    assert quiet.returncode == 0 and quiet.stdout == "Assignment Score: 0 / 0\n" and quiet.stderr == ""


@pytest.mark.parametrize("annotation", [None, "", "not: [valid", "rule:\n  kind: literal"])
def test_invalid_or_missing_annotations_are_skipped_with_warning(tmp_path: Path, annotation: str | None) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx", annotations={("Sheet1", "A1"): annotation})
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    result = run_cli(answer, submission)
    assert result.stdout == "Assignment Score: 0 / 0\n"
    assert result.stderr
    quiet = run_cli(answer, submission, "--quiet-warnings")
    assert quiet.returncode == 0 and quiet.stderr == ""


def test_unknown_rule_kind_is_skipped_with_warning(tmp_path: Path) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx", annotations={("Sheet1", "A1"): "rule:\n  kind: formula\n  points: 9"})
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    result = run_cli(answer, submission)
    assert result.stdout == "Assignment Score: 0 / 0\n"
    assert result.stderr


def test_missing_files_are_fatal_and_name_the_missing_path(tmp_path: Path) -> None:
    answer = tmp_path / "missing-answer.xlsx"
    submission = tmp_path / "missing-submission.xlsx"
    existing_answer = make_answer_key(tmp_path / "existing-answer.xlsx")
    existing_submission = make_submission(tmp_path / "existing-submission.xlsx", {"Sheet1": {"A1": 1}})
    missing_answer_result = run_cli(answer, existing_submission)
    assert missing_answer_result.returncode == 1 and str(answer) in missing_answer_result.stderr and missing_answer_result.stdout == ""
    missing_submission_result = run_cli(existing_answer, submission)
    assert missing_submission_result.returncode == 1 and str(submission) in missing_submission_result.stderr and missing_submission_result.stdout == ""


def test_missing_manifest_is_fatal(tmp_path: Path) -> None:
    answer = make_workbook(tmp_path / "answer.xlsx", {"Sheet1": {"A1": 1}})
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    result = run_cli(answer, submission)
    assert result.returncode == 1
    assert "EvalManifest" in result.stderr and result.stdout == ""


def test_invalid_coordinate_is_fatal_with_or_without_debug(tmp_path: Path) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx", checklist_rows=[["not-a-cell"]])
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    normal = run_cli(answer, submission)
    assert normal.returncode == 1 and normal.stdout == "" and "Error:" in normal.stderr
    debug = run_cli(answer, submission, "--debug")
    assert debug.returncode == 1 and debug.stdout == "" and "Traceback" in debug.stderr


def test_short_options_match_long_options(tmp_path: Path) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx")
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 0}})
    long_result = run_cli(answer, submission, "--log")
    short_result = run_cli(answer, submission, "-l")
    assert (short_result.returncode, short_result.stdout, short_result.stderr) == (
        long_result.returncode,
        long_result.stdout,
        long_result.stderr,
    )


def test_debug_has_no_effect_when_grading_succeeds(tmp_path: Path) -> None:
    answer = make_answer_key(tmp_path / "answer.xlsx")
    submission = make_submission(tmp_path / "submission.xlsx", {"Sheet1": {"A1": 1}})
    normal = run_cli(answer, submission)
    debug = run_cli(answer, submission, "-d")
    assert (debug.returncode, debug.stdout, debug.stderr) == (normal.returncode, normal.stdout, normal.stderr)
